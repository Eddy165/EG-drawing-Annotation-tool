
import json
import os
from typing import List, Dict, Any
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.platypus import Table, TableStyle
import openpyxl
from openpyxl.styles import Font, PatternFill

class ReportGenerator:
    def __init__(self, output_dir: str = "backend/reports"):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)

    def generate_json(self, data: Dict[str, Any], filename: str) -> str:
        filepath = os.path.join(self.output_dir, f"{filename}.json")
        with open(filepath, 'w') as f:
            json.dump(data, f, indent=4)
        return filepath

    def generate_xlsx(self, data: Dict[str, Any], filename: str) -> str:
        filepath = os.path.join(self.output_dir, f"{filename}.xlsx")
        wb = openpyxl.Workbook()
        
        # Sheet 1: Dimensions
        ws1 = wb.active
        ws1.title = "Dimensions"
        headers = ["Balloon No", "View", "Type", "Value", "Unit", "X", "Y"]
        ws1.append(headers)
        
        for b in data.get("balloons", []):
            if b.get("section") == "drawing":
                ws1.append([
                    b.get("balloon_no"),
                    b.get("view_id"),
                    b.get("type"),
                    b.get("value"),
                    b.get("unit"),
                    b.get("position", [0,0])[0],
                    b.get("position", [0,0])[1]
                ])

        # Style Header
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="4F81BD")
        for cell in ws1[1]:
            cell.font = header_font
            cell.fill = header_fill

        # Sheet 2: Notes
        ws2 = wb.create_sheet("Notes")
        ws2.append(["Balloon No", "Note Text"])
        for n in data.get("notes", []):
             ws2.append([n.get("balloon_no"), n.get("text")])
             
        # Sheet 3: BOM
        ws3 = wb.create_sheet("BOM")
        ws3.append(["Balloon No", "Item No", "Description", "Qty", "Material"])
        for item in data.get("bom", []):
            ws3.append([
                item.get("balloon_no"),
                item.get("item_no"),
                item.get("description"),
                item.get("qty"),
                item.get("material")
            ])
            
        wb.save(filepath)
        return filepath

    def generate_pdf(self, data: Dict[str, Any], filename: str) -> str:
        filepath = os.path.join(self.output_dir, f"{filename}.pdf")
        c = canvas.Canvas(filepath, pagesize=A4)
        width, height = A4
        
        # Cover Page
        c.setFont("Helvetica-Bold", 24)
        c.drawString(100, height - 100, "Inspection Report")
        
        c.setFont("Helvetica", 12)
        meta = data.get("metadata", {})
        y = height - 150
        c.drawString(100, y, f"Drawing: {meta.get('drawing_name', 'N/A')}")
        c.drawString(100, y - 20, f"Inspector: {meta.get('inspector', 'N/A')}")
        c.drawString(100, y - 40, f"Date: {meta.get('date', 'N/A')}")
        c.drawString(100, y - 60, f"Total Balloons: {meta.get('total_balloons', 0)}")
        
        c.showPage()
        
        # Table of Balloons (Simple implementation using ReportLab Table)
        # In a real app we'd handle pagination carefully.
        # This is a basic MVP implementation.
        
        data_rows = [["No", "Type", "Value", "Unit"]]
        for b in data.get("balloons", []):
            if b.get("section") == "drawing":
                data_rows.append([
                    str(b.get("balloon_no")),
                    b.get("type", ""),
                    b.get("value", ""),
                    b.get("unit", "")
                ])
                
        table = Table(data_rows, colWidths=[50, 100, 100, 50])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        
        table.wrapOn(c, width, height)
        table.drawOn(c, 50, height - 100 - (len(data_rows) * 20)) # Very rough positioning
        
        c.save()
        return filepath
